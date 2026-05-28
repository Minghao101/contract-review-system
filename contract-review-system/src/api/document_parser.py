"""
文档解析服务 - 支持PDF、DOCX、TXT文件解析
"""
from typing import Any, Dict, Optional
from pathlib import Path
import logging

logger = logging.getLogger(__name__)


class DocumentParser:
    """
    文档解析器

    支持PDF、DOCX、TXT文件的文本提取
    """

    SUPPORTED_FORMATS = {".pdf", ".docx", ".txt"}

    @staticmethod
    async def parse_file(
        file_bytes: bytes,
        filename: str,
    ) -> Dict[str, Any]:
        """
        解析上传的文件

        Args:
            file_bytes: 文件字节内容
            filename: 文件名

        Returns:
            解析结果，包含full_text字段
        """
        suffix = Path(filename).suffix.lower()

        if suffix not in DocumentParser.SUPPORTED_FORMATS:
            return {
                "error": f"不支持的文件格式: {suffix}，支持: {', '.join(DocumentParser.SUPPORTED_FORMATS)}"
            }

        try:
            if suffix == ".pdf":
                return await DocumentParser._parse_pdf(file_bytes, filename)
            elif suffix == ".docx":
                return await DocumentParser._parse_docx(file_bytes, filename)
            elif suffix == ".txt":
                return await DocumentParser._parse_txt(file_bytes, filename)
        except Exception as e:
            logger.error(f"文件解析失败: {e}")
            return {"error": f"文件解析失败: {str(e)}"}

    @staticmethod
    async def _parse_pdf(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        """解析PDF文件"""
        try:
            from pypdf import PdfReader
            import io

            reader = PdfReader(io.BytesIO(file_bytes))
            pages = []
            full_text = []

            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                pages.append({
                    "page_number": i + 1,
                    "char_count": len(page_text),
                })
                full_text.append(page_text)

            text = "\n".join(full_text)

            return {
                "success": True,
                "filename": filename,
                "file_type": "pdf",
                "page_count": len(pages),
                "total_chars": len(text),
                "full_text": text,
            }

        except ImportError:
            return {"error": "请安装pypdf: pip install pypdf"}

    @staticmethod
    async def _parse_docx(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        """解析DOCX文件"""
        try:
            from docx import Document
            import io

            doc = Document(io.BytesIO(file_bytes))
            full_text = []

            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)

            # 提取表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells if cell.text.strip())
                    if row_text:
                        full_text.append(row_text)

            text = "\n".join(full_text)

            return {
                "success": True,
                "filename": filename,
                "file_type": "docx",
                "paragraph_count": len(doc.paragraphs),
                "total_chars": len(text),
                "full_text": text,
            }

        except ImportError:
            return {"error": "请安装python-docx: pip install python-docx"}

    @staticmethod
    async def _parse_txt(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        """解析TXT文件"""
        # 尝试多种编码
        encodings = ["utf-8", "gbk", "gb2312", "latin-1"]
        text = None

        for encoding in encodings:
            try:
                text = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if text is None:
            return {"error": "无法识别文件编码"}

        return {
            "success": True,
            "filename": filename,
            "file_type": "txt",
            "total_chars": len(text),
            "full_text": text,
        }
