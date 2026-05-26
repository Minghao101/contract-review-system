"""
Word文档解析Skill - 解析Docx文件
"""
from typing import Any, Dict
import os
import logging

from ..base_skill import BaseSkill

logger = logging.getLogger(__name__)


class DocxParserSkill(BaseSkill):
    """
    Word文档解析Skill

    解析Docx文件并提取文本内容
    """

    def __init__(self):
        super().__init__(
            skill_id="docx_parser",
            name="Word文档解析",
            description="解析Docx文件并提取文本内容",
        )

    async def execute(self, **kwargs) -> Dict[str, Any]:
        """
        执行Word文档解析

        Args:
            **kwargs:
                - file_path: Docx文件路径
                - file_bytes: Docx文件字节

        Returns:
            提取的文本内容
        """
        file_path = kwargs.get("file_path")
        file_bytes = kwargs.get("file_bytes")

        if not file_path and not file_bytes:
            return {"error": "请提供file_path或file_bytes"}

        try:
            if file_path:
                return self._parse_from_path(file_path)
            else:
                return self._parse_from_bytes(file_bytes)
        except Exception as e:
            logger.error(f"Word文档解析失败: {e}")
            return {"error": str(e)}

    def _parse_from_path(self, file_path: str) -> Dict[str, Any]:
        """从文件路径解析Word文档"""
        if not os.path.exists(file_path):
            return {"error": f"文件不存在: {file_path}"}

        with open(file_path, "rb") as f:
            file_bytes = f.read()

        return self._parse_from_bytes(file_bytes, file_path)

    def _parse_from_bytes(self, file_bytes: bytes, file_path: str = "unknown") -> Dict[str, Any]:
        """从字节解析Word文档"""
        try:
            from docx import Document
            import io

            doc = Document(io.BytesIO(file_bytes))

            paragraphs = []
            full_text = []

            for i, para in enumerate(doc.paragraphs):
                if para.text.strip():
                    paragraphs.append({
                        "index": i,
                        "text": para.text,
                        "style": para.style.name if para.style else None,
                    })
                    full_text.append(para.text)

            # 提取表格内容
            tables = []
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                tables.append(table_data)

            return {
                "file_path": file_path,
                "paragraph_count": len(paragraphs),
                "table_count": len(tables),
                "paragraphs": paragraphs,
                "tables": tables,
                "full_text": "\n".join(full_text),
                "total_chars": sum(len(p["text"]) for p in paragraphs),
            }

        except ImportError:
            return {"error": "请安装python-docx: pip install python-docx"}
        except Exception as e:
            return {"error": f"Word文档解析失败: {str(e)}"}
